"""
FaceAttend AI Pro - Excel Automation Manager
Generates monthly workbooks, dynamic sheets (January-2027 format), and summary statistics.
"""

import os
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import config
from core.database import DatabaseManager
from core.logger import log_info, log_error

class ExcelManager:
    def __init__(self, db: DatabaseManager, excel_path: Path = config.EXCEL_PATH):
        self.db = db
        self.excel_path = excel_path

    def _get_or_create_workbook(self) -> openpyxl.Workbook:
        if self.excel_path.exists():
            try:
                return openpyxl.load_workbook(self.excel_path)
            except Exception as e:
                log_error(f"Error loading Excel workbook, creating new one: {e}")

        wb = openpyxl.Workbook()
        # Remove default sheet if brand new
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])
        return wb

    def generate_monthly_sheet(self, date_obj: datetime = None) -> str:
        """
        Generate or update monthly attendance sheet (e.g., 'January-2027').
        Populates headers (Roll No, Name, Department, 01..31, Present, Absent, Late, %).
        """
        if not date_obj:
            date_obj = datetime.now()

        month_sheet_name = date_obj.strftime("%B-%Y")
        year = date_obj.year
        month = date_obj.month

        import calendar
        days_in_month = calendar.monthrange(year, month)[1]

        wb = self._get_or_create_workbook()

        if month_sheet_name in wb.sheetnames:
            ws = wb[month_sheet_name]
        else:
            ws = wb.create_sheet(title=month_sheet_name)

        # Build Headers
        headers = ["Roll No", "Name", "Department"] + [f"{d:02d}" for d in range(1, days_in_month + 1)] + ["Present", "Absent", "Late", "%"]

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        # Write Headers if not present or empty sheet
        if ws.max_row <= 1:
            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = thin_border
            ws.row_dimensions[1].height = 25

        # Fetch student list and attendance records for month
        students = self.db.get_all_students()

        # Build date prefix for month: YYYY-MM
        month_prefix = f"{year:04d}-{month:02d}"

        # Write student rows
        current_row = 2
        for s in students:
            # Check if student already in sheet row
            existing_row = None
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(row=r, column=1).value) == str(s['roll_no']):
                    existing_row = r
                    break

            r_idx = existing_row if existing_row else current_row
            ws.cell(row=r_idx, column=1, value=s['roll_no']).alignment = left_align
            ws.cell(row=r_idx, column=2, value=s['name']).alignment = left_align
            ws.cell(row=r_idx, column=3, value=s['department']).alignment = left_align

            present_cnt = 0
            late_cnt = 0
            absent_cnt = 0

            # Fill days
            for d in range(1, days_in_month + 1):
                date_str = f"{month_prefix}-{d:02d}"
                col_idx = 3 + d  # 1:Roll, 2:Name, 3:Dept, 4:01, etc.

                # Query database attendance for this student on date_str
                with self.db.get_connection() as conn:
                    cursor = conn.cursor()
                    cursor.execute("SELECT status FROM attendance WHERE student_id = ? AND date = ?;", (s['student_id'], date_str))
                    row = cursor.fetchone()

                status = row['status'] if row else ""
                cell = ws.cell(row=r_idx, column=col_idx, value=status)
                cell.alignment = center_align
                cell.border = thin_border

                if status == 'P':
                    present_cnt += 1
                    cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")  # Light green
                    cell.font = Font(color="065F46", bold=True)
                elif status == 'L':
                    late_cnt += 1
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Light yellow
                    cell.font = Font(color="92400E", bold=True)
                else:
                    # Check if date is in past to mark absent
                    today_str = datetime.now().strftime("%Y-%m-%d")
                    if date_str < today_str:
                        absent_cnt += 1
                        cell.value = 'A'
                        cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # Light red
                        cell.font = Font(color="991B1B", bold=True)

            total_days = days_in_month
            effective_present = present_cnt + late_cnt
            pct = round((effective_present / total_days * 100.0), 1) if total_days > 0 else 0.0

            # Write Summary Columns
            p_col = 3 + days_in_month + 1
            a_col = 3 + days_in_month + 2
            l_col = 3 + days_in_month + 3
            pct_col = 3 + days_in_month + 4

            ws.cell(row=r_idx, column=p_col, value=present_cnt).alignment = center_align
            ws.cell(row=r_idx, column=a_col, value=absent_cnt).alignment = center_align
            ws.cell(row=r_idx, column=l_col, value=late_cnt).alignment = center_align
            ws.cell(row=r_idx, column=pct_col, value=f"{pct}%").alignment = center_align

            if not existing_row:
                current_row += 1

        # Auto column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 6)

        os.makedirs(config.ATTENDANCE_DIR, exist_ok=True)
        wb.save(self.excel_path)
        log_info(f"Excel workbook saved with sheet '{month_sheet_name}'.")
        return str(self.excel_path)
